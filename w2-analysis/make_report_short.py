#!/usr/bin/env python3
"""Краткий отчёт (полстраницы, без таблиц). Цифры берутся из tx*.json через analyze.py."""
import json
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from analyze import analyze

SIGS = {n: (net, s) for n, net, s in (l.split() for l in open('signatures.txt'))}
T = {}
for n, (net, _) in SIGS.items():
    T[n] = {}
    analyze(json.load(open(f'{n}.json')), net, T[n])
sp = lambda n: f'{n:,}'.replace(',', ' ')

doc = Document()
st = doc.styles['Normal']
st.font.name, st.font.size = 'Calibri', Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
h = doc.styles['Heading 1']
h.font.color.rgb, h.font.size = RGBColor(0, 0, 0), Pt(14)
rf = h.element.rPr.rFonts
for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
    rf.attrib.pop(qn(a), None)
rf.set(qn('w:ascii'), 'Calibri'); rf.set(qn('w:hAnsi'), 'Calibri')


def explorer_link(p, name):
    net, sig = SIGS[name]
    url = f'https://explorer.solana.com/tx/{sig}' + ('?cluster=devnet' if net == 'devnet' else '')
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement('w:hyperlink'); hl.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u); r.append(rpr)
    t = OxmlElement('w:t'); t.text = 'Explorer'; r.append(t); hl.append(r)
    p.add_run(' ('); p._p.append(hl); p.add_run(').')


def item(name, text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f'{name}: ').bold = True
    p.add_run(text)
    explorer_link(p, name)


t0, t1, t2, t3 = T['tx0'], T['tx1'], T['tx2'], T['tx3']
doc.add_heading('Задание 2. Как на самом деле работает Solana: анализ транзакций', level=1)
doc.add_paragraph('Через JSON-RPC (getTransaction) получил три транзакции из mainnet и свою из недели 1, '
                  'разобрал их скриптом analyze.py: accounts, instructions (включая inner), programs, комиссия, CU, балансы.')
item('tx1', f"перевод SOL — {len(t1['accounts'])} аккаунта, 1 инструкция transfer System Program, "
            f"{t1['cu']} CU, комиссия {sp(t1['fee'])} lamports")
item('tx2', f"перевод 11.765859 USDC — {len(t2['accounts'])} аккаунтов, {t2['n_top']} инструкций "
            f"(Compute Budget, Memo, ATA, SPL Token transferChecked), {sp(t2['cu'])} CU, комиссия {sp(t2['fee'])} lamports")
item('tx3', f"swap через Jupiter v6 (SOL → TTWO → banked через пулы Raydium) — {len(t3['accounts'])} аккаунтов, "
            f"{t3['n_top']} инструкций + {t3['n_inner']} inner, {len(t3['programs'])} программ, "
            f"{sp(t3['cu'])} CU, комиссия {sp(t3['fee'])} lamports")
item('tx0', f"мой перевод 0.1 SOL на Phantom в devnet — та же структура, что у tx1")
doc.add_heading('Вывод', level=1)
doc.add_paragraph(
    'Транзакция — это набор инструкций к программам и заранее объявленный список аккаунтов с правами signer/writable. '
    f"Swap сложнее перевода примерно в {round(t3['cu'] / t1['cu'])} раз по compute units, потому что Jupiter через CPI вызывает "
    f"несколько программ, но комиссия выше лишь в {t3['fee'] / t1['fee']:.1f} раза: базовая часть одинакова (5 000 lamports за подпись), "
    'растёт только priority fee. Реальное движение средств лучше всего видно по pre/post балансам.')
doc.save('report_week2.docx')
print('saved report_week2.docx')
